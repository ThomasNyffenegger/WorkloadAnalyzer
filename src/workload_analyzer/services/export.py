import csv
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook

from workload_analyzer.core.rounding import round_seconds
from workload_analyzer.db.repository import Repository


def _ts_to_iso(ts: int) -> str:
    return datetime.fromtimestamp(ts, tz=timezone.utc).astimezone().isoformat(timespec="seconds")


def _gather_rows(repo: Repository, from_ts: int, to_ts: int, rounding_minutes: int):
    """Yield (start, end, duration_min, role, category, source, comment) tuples."""
    cats = {c.id: c for c in repo.list_categories()}
    roles = {r.id: r for r in repo.list_roles()}
    for entry in repo.list_entries_between(from_ts, to_ts):
        if entry.end_ts is None:
            continue
        cat = cats.get(entry.category_id)
        if cat is None:
            continue
        role = roles.get(cat.role_id)
        role_name = role.name if role else ""
        duration_s = round_seconds(entry.duration_seconds(), rounding_minutes)
        if duration_s == 0 and rounding_minutes != 0:
            continue
        yield (
            _ts_to_iso(entry.start_ts),
            _ts_to_iso(entry.end_ts),
            duration_s // 60,
            role_name,
            cat.name,
            entry.source.value,
            entry.comment or "",
        )


def export_csv(repo: Repository, from_ts: int, to_ts: int, out_path: Path, rounding_minutes: int = 0) -> None:
    with out_path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["start", "end", "duration_minutes", "role", "category", "source", "comment"])
        for row in _gather_rows(repo, from_ts, to_ts, rounding_minutes):
            writer.writerow(row)


def export_xlsx(repo: Repository, from_ts: int, to_ts: int, rounding_minutes: int, out_path: Path) -> None:
    rows = list(_gather_rows(repo, from_ts, to_ts, rounding_minutes))

    wb = Workbook()
    # Summary first
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Role", "Category", "Total minutes"])
    totals: dict[tuple[str, str], int] = defaultdict(int)
    for start, end, dur, role, cat, src, comment in rows:
        totals[(role, cat)] += dur
    for (role, cat), dur in sorted(totals.items()):
        summary.append([role, cat, dur])

    # Details
    details = wb.create_sheet("Details")
    details.append(["start", "end", "duration_minutes", "role", "category", "source", "comment"])
    for row in rows:
        details.append(list(row))

    wb.save(out_path)


def export_xlsx_pivot(
    repo: Repository,
    from_ts: int,
    to_ts: int,
    rounding_minutes: int,
    out_path: Path,
) -> None:
    """Export a Tage × Kategorien pivot table to a single-sheet XLSX file.

    Rows = calendar days in [from_ts, to_ts).
    Columns = categories that have at least one entry in the range (alphabetical).
    Cells = total hours (float, 1 decimal).
    Last column = daily total. Last row = column totals.
    """
    from datetime import date, timedelta
    from openpyxl.styles import Font, PatternFill

    cats = {c.id: c for c in repo.list_categories()}
    entries = [e for e in repo.list_entries_between(from_ts, to_ts) if e.end_ts is not None]

    # Date range: one row per calendar day
    from_date = datetime.fromtimestamp(from_ts, tz=timezone.utc).astimezone().date()
    to_date = datetime.fromtimestamp(max(from_ts, to_ts - 1), tz=timezone.utc).astimezone().date()
    days: list[date] = []
    d = from_date
    while d <= to_date:
        days.append(d)
        d += timedelta(days=1)

    # Categories that appear in entries, sorted alphabetically
    used_cat_ids = sorted(
        {e.category_id for e in entries if e.category_id in cats},
        key=lambda cid: cats[cid].name,
    )

    # Aggregate: (date, cat_id) -> hours
    totals: dict[tuple[date, int], float] = {}
    for entry in entries:
        cat = cats.get(entry.category_id)
        if cat is None:
            continue
        dur_s = round_seconds(entry.duration_seconds(), rounding_minutes)
        entry_date = datetime.fromtimestamp(entry.start_ts, tz=timezone.utc).astimezone().date()
        key = (entry_date, entry.category_id)
        totals[key] = totals.get(key, 0.0) + dur_s / 3600.0

    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot"

    total_col = len(used_cat_ids) + 2  # col1=Datum, col2..N=categories, colN+1=Total

    # Header row
    ws.cell(row=1, column=1, value="Datum").font = Font(bold=True)
    for col_idx, cat_id in enumerate(used_cat_ids, start=2):
        cat = cats[cat_id]
        cell = ws.cell(row=1, column=col_idx, value=cat.name)
        cell.font = Font(bold=True)
        if cat.color and len(cat.color.lstrip("#")) == 6:
            hex_color = cat.color.lstrip("#")
            cell.fill = PatternFill(fill_type="solid", fgColor=hex_color)
    ws.cell(row=1, column=total_col, value="Total").font = Font(bold=True)

    # Data rows (one per day)
    for row_idx, day in enumerate(days, start=2):
        ws.cell(row=row_idx, column=1, value=day.strftime("%d.%m.%y"))
        row_total = 0.0
        for col_idx, cat_id in enumerate(used_cat_ids, start=2):
            hours = round(totals.get((day, cat_id), 0.0), 1)
            ws.cell(row=row_idx, column=col_idx, value=hours)
            row_total += hours
        ws.cell(row=row_idx, column=total_col, value=round(row_total, 1))

    # Totals row
    total_row = len(days) + 2
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    grand_total = 0.0
    for col_idx, cat_id in enumerate(used_cat_ids, start=2):
        col_total = round(sum(round(totals.get((day, cat_id), 0.0), 1) for day in days), 1)
        ws.cell(row=total_row, column=col_idx, value=col_total).font = Font(bold=True)
        grand_total += col_total
    ws.cell(row=total_row, column=total_col, value=round(grand_total, 1)).font = Font(bold=True)

    wb.save(out_path)
