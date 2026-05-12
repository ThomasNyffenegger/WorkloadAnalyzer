import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from workload_analyzer.db.connection import connect
from workload_analyzer.db.repository import Repository
from workload_analyzer.models import EntrySource
from workload_analyzer.services.export import export_csv, export_xlsx


@pytest.fixture
def repo_with_entries(tmp_db_path):
    conn = connect(tmp_db_path)
    repo = Repository(conn)
    rid = repo.create_role("Entwicklung")
    cid = repo.create_category(name="Frontend", color="#0", role_id=rid)
    repo.insert_closed_entry(category_id=cid, start_ts=1700000000, end_ts=1700003600, source=EntrySource.MANUAL)
    repo.insert_closed_entry(category_id=cid, start_ts=1700003600, end_ts=1700007200, source=EntrySource.MANUAL, comment="note")
    yield repo
    conn.close()


def test_export_csv_writes_rows(repo_with_entries, tmp_path: Path):
    out = tmp_path / "out.csv"
    export_csv(repo_with_entries, from_ts=0, to_ts=9999999999, out_path=out)
    with out.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == ["start", "end", "duration_minutes", "role", "category", "source", "comment"]
    assert len(rows) == 3
    assert rows[1][3] == "Entwicklung"
    assert rows[1][4] == "Frontend"


def test_export_xlsx_has_summary_and_details_sheets(repo_with_entries, tmp_path: Path):
    out = tmp_path / "out.xlsx"
    export_xlsx(repo_with_entries, from_ts=0, to_ts=9999999999, rounding_minutes=0, out_path=out)
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "Details" in wb.sheetnames
    details = wb["Details"]
    headers = [c.value for c in details[1]]
    assert "category" in headers
    assert "duration_minutes" in headers


def test_export_xlsx_summary_aggregates_per_category(repo_with_entries, tmp_path: Path):
    out = tmp_path / "out.xlsx"
    export_xlsx(repo_with_entries, from_ts=0, to_ts=9999999999, rounding_minutes=0, out_path=out)
    wb = load_workbook(out)
    summary = wb["Summary"]
    rows = list(summary.iter_rows(values_only=True))
    # Expect header + one data row for "Frontend" with 120 minutes total
    data_rows = [r for r in rows[1:] if r[0] is not None]
    assert len(data_rows) == 1
    role, category, total_minutes = data_rows[0][0], data_rows[0][1], data_rows[0][2]
    assert role == "Entwicklung"
    assert category == "Frontend"
    assert total_minutes == 120
