"""Tests for export_xlsx_pivot()."""
import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from workload_analyzer.db.connection import connect
from workload_analyzer.db.repository import Repository
from workload_analyzer.models import EntrySource
from workload_analyzer.services.export import export_xlsx_pivot


def make_repo(tmp_db_path):
    conn = connect(tmp_db_path)
    repo = Repository(conn)
    role_id = repo.create_role("Dev")
    cat1 = repo.create_category("Coding", "#3388cc", role_id)
    cat2 = repo.create_category("Meetings", "#ff8800", role_id)
    return repo, cat1, cat2


def day_ts(date_str: str) -> int:
    """Return unix timestamp for midnight local time of the given date."""
    d = datetime.datetime.strptime(date_str, "%Y-%m-%d")
    return int(d.astimezone(datetime.timezone.utc).timestamp())


def test_pivot_creates_xlsx_file(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    assert out.exists()


def test_pivot_sheet_named_pivot(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    wb = load_workbook(out)
    assert "Pivot" in wb.sheetnames


def test_pivot_header_contains_datum_category_total(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    start = day_ts("2026-05-01") + 3600
    repo.start_entry(cat1, start, EntrySource.MANUAL)
    repo.close_entry(repo.get_open_entry().id, start + 3600)
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    ws = load_workbook(out)["Pivot"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert headers[0] == "Datum"
    assert "Coding" in headers
    assert headers[-1] == "Total"


def test_pivot_hours_computed_correctly(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    start = day_ts("2026-05-01") + 3600
    repo.start_entry(cat1, start, EntrySource.MANUAL)
    repo.close_entry(repo.get_open_entry().id, start + 5400)  # 1.5 hours
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    ws = load_workbook(out)["Pivot"]
    coding_col = next(
        c for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value == "Coding"
    )
    assert ws.cell(row=2, column=coding_col).value == 1.5


def test_pivot_totals_row_is_last_and_labelled(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    start = day_ts("2026-05-01") + 3600
    repo.start_entry(cat1, start, EntrySource.MANUAL)
    repo.close_entry(repo.get_open_entry().id, start + 3600)
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    ws = load_workbook(out)["Pivot"]
    assert ws.cell(row=ws.max_row, column=1).value == "Total"


def test_pivot_days_with_no_entries_show_zero(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    # Entry only on day 1 of a 2-day range
    start = day_ts("2026-05-01") + 3600
    repo.start_entry(cat1, start, EntrySource.MANUAL)
    repo.close_entry(repo.get_open_entry().id, start + 3600)
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-03"), 0, out)
    ws = load_workbook(out)["Pivot"]
    coding_col = next(
        c for c in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=c).value == "Coding"
    )
    # Row 3 = second day (2026-05-02) — no entries, should be 0.0
    assert ws.cell(row=3, column=coding_col).value == 0.0


def test_pivot_empty_range_produces_valid_file(tmp_db_path, tmp_path):
    repo, cat1, cat2 = make_repo(tmp_db_path)
    # No entries at all
    out = tmp_path / "pivot.xlsx"
    export_xlsx_pivot(repo, day_ts("2026-05-01"), day_ts("2026-05-02"), 0, out)
    assert out.exists()
    wb = load_workbook(out)
    assert "Pivot" in wb.sheetnames
